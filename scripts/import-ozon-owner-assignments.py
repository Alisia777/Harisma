#!/usr/bin/env python3
"""Import Ozon-only owner assignments from a three-column XLSX file."""

from __future__ import annotations

import argparse
import gzip
import json
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook


SOURCE_NOTE = "ozon-owner-assignments-2026-06-30"
OWNER_ALIASES = {
    "\u0414\u0430\u0448\u0430": "\u041c\u043e\u043b\u043e\u0434\u044f\u043a\u043e\u0432\u0430 \u0414\u0430\u0440\u0438\u044f",
    "\u0414\u0430\u0440\u044c\u044f": "\u041c\u043e\u043b\u043e\u0434\u044f\u043a\u043e\u0432\u0430 \u0414\u0430\u0440\u0438\u044f",
    "\u0414\u0430\u0440\u0438\u044f": "\u041c\u043e\u043b\u043e\u0434\u044f\u043a\u043e\u0432\u0430 \u0414\u0430\u0440\u0438\u044f",
    "\u0410\u0440\u0442\u0435\u043c": "\u041f\u0438\u0442\u0430\u0439\u043a\u0438\u043d \u0410\u0440\u0442\u0451\u043c",
    "\u0410\u0440\u0442\u0451\u043c": "\u041f\u0438\u0442\u0430\u0439\u043a\u0438\u043d \u0410\u0440\u0442\u0451\u043c",
}
PLATFORM_OWNER_KEYS = {
    "ownerByPlatform",
    "ownersByPlatform",
    "platformOwners",
}
TARGET_FILES = [
    "data/sku_matrix.json",
    "data/dashboard.json",
    "data/prices.json",
    "data/repricer.json",
    "data/order_procurement_ozon.json",
]
OWNER_DISPLAY_PLATFORMS = ["wb", "ozon", "ym", "ga", "letu", "megamarket", "samokat", "mm", "magnit"]


def normalize_article(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"[^\w\-\u0400-\u04ff]+", "_", text, flags=re.UNICODE)
    text = re.sub(r"_+", "_", text)
    return text.strip("_")


def normalize_token(value: Any) -> str:
    return re.sub(r"[^\w\u0400-\u04ff]+", "", str(value or "").strip().lower(), flags=re.UNICODE)


def clean_owner(value: Any) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    return OWNER_ALIASES.get(text, text)


def read_json(path: Path, fallback: Any = None) -> Any:
    if not path.exists():
        return fallback
    return json.loads(path.read_text(encoding="utf-8-sig"))


def write_json(path: Path, payload: Any, compact: bool = False) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if compact:
        path.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    else:
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_assignments(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows: list[dict[str, str]] = []
    seen: dict[str, str] = {}
    for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        article = str(row[0] or "").strip() if len(row) > 0 else ""
        sku_id = str(row[1] or "").strip() if len(row) > 1 else ""
        owner = clean_owner(row[2] if len(row) > 2 else "")
        if not article and not sku_id and not owner:
            continue
        if not article or not owner:
            continue
        article_key = normalize_article(article)
        previous = seen.get(article_key)
        if previous and previous != owner:
            raise ValueError(f"Conflicting owner for {article}: {previous} vs {owner}")
        seen[article_key] = owner
        rows.append({
            "row": row_index,
            "article": article,
            "article_key": article_key,
            "sku_id": sku_id,
            "sku_token": normalize_token(sku_id),
            "owner": owner,
        })
    return rows


def build_assignment_indexes(assignments: list[dict[str, str]]) -> tuple[dict[str, str], dict[str, str]]:
    by_article: dict[str, str] = {}
    by_sku: dict[str, str] = {}
    for item in assignments:
        by_article[item["article_key"]] = item["owner"]
        if item["sku_token"]:
            by_sku[item["sku_token"]] = item["owner"]
    return by_article, by_sku


def item_tokens(item: dict[str, Any]) -> tuple[list[str], list[str]]:
    article_fields = [
        item.get("articleKey"),
        item.get("article"),
        item.get("sourceArticleKey"),
        item.get("vendorCode"),
        item.get("offerId"),
    ]
    sku_fields = [
        item.get("sku"),
        item.get("skuId"),
        item.get("marketArticleId"),
        item.get("productId"),
        item.get("nmId"),
    ]
    return (
        [normalize_article(value) for value in article_fields if str(value or "").strip()],
        [normalize_token(value) for value in sku_fields if str(value or "").strip()],
    )


def owner_for_item(item: dict[str, Any], by_article: dict[str, str], by_sku: dict[str, str]) -> str:
    article_tokens, sku_tokens = item_tokens(item)
    for token in article_tokens:
        owner = by_article.get(token)
        if owner:
            return owner
    for token in sku_tokens:
        owner = by_sku.get(token)
        if owner:
            return owner
    return ""


def is_ozon_context(item: dict[str, Any], inherited_platform: str = "") -> bool:
    platform = str(item.get("platformKey") or item.get("platform") or item.get("key") or inherited_platform or "").lower()
    return platform in {"ozon", "oz"}


def ensure_owner_object(sku: dict[str, Any]) -> None:
    owner = sku.get("owner")
    if not isinstance(owner, dict):
        sku["owner"] = {"name": str(owner or "").strip(), "source": "", "byPlatform": {}}
    if not isinstance(sku["owner"].get("byPlatform"), dict):
        sku["owner"]["byPlatform"] = {}
    if not isinstance(sku.get("ownersByPlatform"), dict):
        sku["ownersByPlatform"] = {}


def apply_sku_owners(root: Path, by_article: dict[str, str], by_sku: dict[str, str], source_label: str) -> dict[str, Any]:
    path = root / "data/skus.json"
    skus = read_json(path, [])
    touched = 0
    changes = Counter()
    unmatched = []
    assignment_hits = set()
    for sku in skus:
        if not isinstance(sku, dict):
            continue
        owner = owner_for_item(sku, by_article, by_sku)
        if not owner:
            unmatched.append(sku.get("articleKey") or sku.get("article") or "")
            continue
        assignment_hits.add(normalize_article(sku.get("articleKey") or sku.get("article")))
        ensure_owner_object(sku)
        old_owner = sku["ownersByPlatform"].get("ozon") or sku["owner"]["byPlatform"].get("ozon") or ""
        if old_owner != owner:
            touched += 1
            changes[f"{old_owner or '<empty>'} -> {owner}"] += 1
            sku["owner"]["source"] = f"{SOURCE_NOTE} - {source_label}"
        sku["owner"]["byPlatform"]["ozon"] = owner
        sku["ownersByPlatform"]["ozon"] = owner
    write_json(path, skus)
    return {
        "file": str(path.relative_to(root)),
        "updatedSkus": touched,
        "matchedSkus": len(assignment_hits),
        "ownerChanges": dict(changes),
        "unmatchedPortalSkus": len([item for item in unmatched if item]),
    }


def apply_ozon_owner_to_item(item: dict[str, Any], owner: str) -> bool:
    changed = False
    if item.get("owner") != owner and is_ozon_context(item):
        item["owner"] = owner
        changed = True
    for key in PLATFORM_OWNER_KEYS:
        value = item.get(key)
        if isinstance(value, dict) and value.get("ozon") != owner:
            value["ozon"] = owner
            changed = True
    owner_obj = item.get("owner")
    if isinstance(owner_obj, dict):
        by_platform = owner_obj.setdefault("byPlatform", {})
        if isinstance(by_platform, dict) and by_platform.get("ozon") != owner:
            by_platform["ozon"] = owner
            changed = True
    if "owner_ozon" in item and item.get("owner_ozon") != owner:
        item["owner_ozon"] = owner
        changed = True
    platform_owners = item.get("platformOwners")
    if isinstance(platform_owners, dict) and "owner" in item:
        next_owner = owner_display_text(item.get("productOwner"), platform_owners)
        if next_owner and item.get("owner") != next_owner:
            item["owner"] = next_owner
            changed = True
    return changed


def owner_display_text(product_owner: Any, platform_owners: dict[str, Any]) -> str:
    product_owner = str(product_owner or "").strip()
    entries = []
    for platform in OWNER_DISPLAY_PLATFORMS:
        owner = str(platform_owners.get(platform) or "").strip()
        if owner:
            entries.append((platform, owner))
    for platform, owner in platform_owners.items():
        platform = str(platform or "").strip()
        owner = str(owner or "").strip()
        if owner and platform and all(existing_platform != platform for existing_platform, _ in entries):
            entries.append((platform, owner))
    if not entries:
        return ""
    unique_owners = {owner for _, owner in entries}
    if len(unique_owners) == 1 and (not product_owner or product_owner in unique_owners):
        return entries[0][1]
    return "; ".join(f"{platform}: {owner}" for platform, owner in entries)


def walk_payload(value: Any, by_article: dict[str, str], by_sku: dict[str, str], inherited_platform: str = "") -> int:
    updated = 0
    if isinstance(value, list):
        for item in value:
            updated += walk_payload(item, by_article, by_sku, inherited_platform)
        return updated
    if not isinstance(value, dict):
        return 0

    platform = str(value.get("platformKey") or value.get("platform") or value.get("key") or inherited_platform or "").lower()
    owner = owner_for_item(value, by_article, by_sku)
    has_platform_owner = any(isinstance(value.get(key), dict) for key in PLATFORM_OWNER_KEYS)
    if owner and (is_ozon_context(value, inherited_platform) or has_platform_owner or "owner_ozon" in value):
        if apply_ozon_owner_to_item(value, owner):
            updated += 1

    for key, child in value.items():
        next_platform = platform
        if key == "ozon":
            next_platform = "ozon"
        updated += walk_payload(child, by_article, by_sku, next_platform)
    return updated


def apply_target_files(root: Path, by_article: dict[str, str], by_sku: dict[str, str]) -> list[dict[str, Any]]:
    stats = []
    for relative in TARGET_FILES:
        path = root / relative
        payload = read_json(path, None)
        if payload is None:
            continue
        updated = walk_payload(payload, by_article, by_sku)
        if updated:
            write_json(path, payload, compact="compact" in path.name or "minified" in path.name)
            if relative == "data/order_procurement_ozon.json":
                (root / "data/order_procurement_ozon.json.gz").write_bytes(
                    gzip.compress(path.read_bytes(), compresslevel=9, mtime=0)
                )
        stats.append({"file": relative, "updatedObjects": updated})
    return stats


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", required=True, help="XLSX with columns: article, sku, owner")
    parser.add_argument("--root", default=".", help="Portal repository root")
    parser.add_argument("--source-label", default="", help="Human-readable source label")
    args = parser.parse_args()

    root = Path(args.root).resolve()
    source = Path(args.file).resolve()
    source_label = args.source_label or source.name
    assignments = read_assignments(source)
    by_article, by_sku = build_assignment_indexes(assignments)
    sku_stats = apply_sku_owners(root, by_article, by_sku, source_label)
    file_stats = apply_target_files(root, by_article, by_sku)

    matched_articles = {item["article_key"] for item in assignments if item["article_key"] in by_article}
    report = {
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "sourceFile": source_label,
        "assignmentRows": len(assignments),
        "uniqueArticles": len(matched_articles),
        "ownerCounts": dict(Counter(item["owner"] for item in assignments)),
        "skus": sku_stats,
        "files": file_stats,
    }
    report_path = root / "exports" / "ozon-owner-assignments-2026-06-30.json"
    write_json(report_path, report)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
